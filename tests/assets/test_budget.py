"""Tests for the budget spending profile and the NJ property tax deduction that reads it."""

import io
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

import owlplanner as owl
from owlplanner import budget, spending, tax_state
from owlplanner.export import _compute_estate, budget_spending, plan_metrics
from owlplanner.hfp_io import conditionDebtsAndFixedAssetsDF

THISYEAR = owl.Plan(["A"], ["1960-01-01"], [80], "t", verbose=False).year_n[0]
COLS = ["active", "name", "type", "year", "end", "amount", "rate", "survivor"]


def _line(**kw):
    base = {"active": True, "name": "l", "type": "core", "year": THISYEAR, "end": 0, "amount": 12000.0,
            "rate": 0.0, "survivor": np.nan}
    base.update(kw)
    return base


def _df(*lines):
    return conditionDebtsAndFixedAssetsDF(pd.DataFrame(list(lines), columns=COLS), "Budget")


def _core(amount, **kw):
    return _line(name="core", type="core", amount=amount, **kw)


def _pt(amount=20000.0):
    return _line(name="pt", type="property tax", amount=amount)


def _rent(amount):
    return _line(name="rent", type="rent", amount=amount)


# ---------------------------------------------------------------------------
# budget.evaluate
# ---------------------------------------------------------------------------


class TestEvaluate:
    def test_empty(self):
        b = budget.evaluate(_df(), 5, 5, 60, THISYEAR)
        assert b.amounts_ln.shape == (0, 5) and np.all(b.total_n == 0)

    def test_constant_through_plan_end(self):
        b = budget.evaluate(_df(_rent(24000.0)), 5, 5, 60, THISYEAR)
        assert np.all(b.total_n == 24000.0) and np.all(b.by_type("rent") == 24000.0)

    def test_end_year_is_inclusive(self):
        b = budget.evaluate(_df(_line(end=THISYEAR + 1)), 5, 5, 60, THISYEAR)
        assert list(b.total_n) == [12000.0, 12000.0, 0, 0, 0]

    def test_end_negative_counts_from_the_plan_end(self):
        b = budget.evaluate(_df(_line(end=-1)), 5, 5, 60, THISYEAR)
        assert list(b.total_n) == [12000.0, 12000.0, 12000.0, 12000.0, 0]

    def test_year_bounds_and_past_start(self):
        b = budget.evaluate(_df(_line(year=THISYEAR + 2)), 5, 5, 60, THISYEAR)
        assert list(b.total_n) == [0, 0, 12000.0, 12000.0, 12000.0]
        b = budget.evaluate(_df(_line(year=THISYEAR - 3, rate=2.0)), 3, 3, 60, THISYEAR)
        assert b.total_n[0] == pytest.approx(12000.0)  # growth counts from the plan start

    def test_inactive_line_is_skipped(self):
        b = budget.evaluate(_df(_line(active=False)), 5, 5, 60, THISYEAR)
        assert np.all(b.total_n == 0) and b.names == ()

    def test_real_growth_no_inflation(self):
        b = budget.evaluate(_df(_line(type="care", amount=10000.0, rate=2.0)), 4, 4, 60, THISYEAR)
        assert b.total_n == pytest.approx([10000.0 * 1.02**k for k in range(4)])

    def test_survivor_defaults(self):
        # Personal lines keep the case's percentage, housing lines are kept in full.
        b = budget.evaluate(_df(_core(10000.0), _rent(5000.0)), 6, 3, 60, THISYEAR)
        assert list(b.by_type("core")) == [10000.0] * 3 + [6000.0] * 3
        assert list(b.by_type("rent")) == [5000.0] * 6

    def test_survivor_explicit(self):
        b = budget.evaluate(_df(_rent(5000.0) | {"survivor": 50.0}), 4, 2, 60, THISYEAR)
        assert list(b.total_n) == [5000.0, 5000.0, 2500.0, 2500.0]
        with pytest.raises(ValueError, match="survivor"):
            budget.evaluate(_df(_rent(5000.0) | {"survivor": 150.0}), 4, 2, 60, THISYEAR)

    def test_profile_is_normalized_to_its_first_year(self):
        b = budget.evaluate(_df(_core(10000.0), _line(name="care", type="care", year=THISYEAR + 2)), 4, 4, 60,
                            THISYEAR)
        assert list(budget.profile(b)) == pytest.approx([1.0, 1.0, 2.2, 2.2])
        with pytest.raises(ValueError, match="first year"):
            budget.profile(budget.evaluate(_df(_line(year=THISYEAR + 1)), 4, 4, 60, THISYEAR))

    def test_lines_left_out(self):
        df = _df(_line(name="term typed as end", end=10), _line(name="after the plan", year=THISYEAR + 50),
                 _line(name="fine"), _line(name="inactive", active=False, end=10))
        assert [n for n, _ in budget.lines_left_out(df, 5, THISYEAR)] == ["term typed as end", "after the plan"]

    def test_one_core_line_is_the_flat_profile(self):
        b = budget.evaluate(_df(_core(50000.0)), 30, 20, 60, THISYEAR)
        assert budget.profile(b) == pytest.approx(spending.gen_spending_profile("flat", 0.6, 20, 30))


# ---------------------------------------------------------------------------
# Plan
# ---------------------------------------------------------------------------

EXACT = {"withMedicare": "None", "withSSTaxability": 0.85}


def _plan(state="NJ", lines=None, taxable=(200, 100), profile="budget", log=None):
    streams = [log] if log is not None else None
    p = owl.Plan(["Joe", "Jane"], ["1964-03-15", "1965-09-15"], [89, 92], "budget", verbose=False,
                 logstreams=streams)
    p.setSpendingProfile(profile, 60)
    p.setAccountBalances(taxable=list(taxable), taxDeferred=[0, 0], taxFree=[0, 0])
    p.setAllocationRatios("individual", generic=np.array([[[60, 40, 0, 0], [60, 40, 0, 0]]] * 2))
    p.setRates("conservative")
    if lines is not None:
        p.houseLists["Budget"] = _df(*lines)
    if state:
        p.setStateTax(state)
    return p


def _bequest(p):
    return _compute_estate(p, p.N_n)[3] / p.gamma_n[p.N_n]


def _bequest_opts():
    return {"noRothConversions": "None", "maxRothConversion": 0, "withMedicare": "None"}


@pytest.mark.parametrize("objective", ["maxSpending", "maxBequest"])
def test_one_core_line_solves_as_the_flat_profile(objective):
    opts = {**EXACT, "bequest": 0} if objective == "maxSpending" else {**EXACT, "netSpending": 60}
    p_f = _plan(state="", profile="flat", taxable=(800, 400))
    p_b = _plan(state="", lines=[_core(60000.0)], taxable=(800, 400))
    for p in (p_f, p_b):
        p.setSocialSecurity([2800, 2200], [70, 70])
        p.solve(objective, options=dict(opts))
        assert p.caseStatus == "solved"
    assert p_b.g_n == pytest.approx(p_f.g_n, rel=1e-6)
    assert _bequest(p_b) == pytest.approx(_bequest(p_f), rel=1e-6, abs=1.0)


def test_rent_in_the_budget_equals_rent_as_big_ticket_items():
    """At fixed spending, rent inside the profile gives the plan of rent outside it (Phase 2's ledger).

    The ledger was checked equal to negative big-ticket items; here the big-ticket series carries the
    rent, inflated as the profile inflates it, and the flat profile carries the core spending.
    """
    p_b = _plan(state="", lines=[_core(50000.0), _rent(18000.0)], taxable=(800, 400))
    p_b.setSocialSecurity([2800, 2200], [70, 70])
    p_b.solve("maxBequest", options=dict(EXACT))  # netSpending from the budget: 68,000
    assert p_b.g_n[0] == pytest.approx(68000.0)

    p_t = _plan(state="", profile="flat", taxable=(800, 400))
    p_t.setSocialSecurity([2800, 2200], [70, 70])
    p_t.Lambda_in[0, :] = -18000.0 * p_b.gamma_n[:-1]
    p_t.Lambda_in[1, :] = 0.0
    p_t.solve("maxBequest", options={**EXACT, "netSpending": 50})
    assert _bequest(p_b) == pytest.approx(_bequest(p_t), rel=1e-6)


def test_max_spending_scales_the_whole_budget():
    car = _line(name="car", type="car", amount=10000.0, end=THISYEAR + 4)
    p = _plan(state="", lines=[_core(50000.0), _rent(18000.0), car], taxable=(800, 400))
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.solve("maxSpending", options={**EXACT, "bequest": 0})
    real = p.g_n / p.gamma_n[:-1]
    assert real / real[0] == pytest.approx(p.budget.total_n / p.budget.total_n[0], rel=1e-6)
    lines, housing = budget_spending(p)
    assert sum(lines.values()) == pytest.approx(p.g_n, rel=1e-9)
    assert housing == pytest.approx(lines["rent"])


def test_explicit_net_spending_scales_the_budget():
    p = _plan(state="", lines=[_core(50000.0), _rent(18000.0)], taxable=(800, 400))
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.solve("maxBequest", options={**EXACT, "netSpending": 34})
    assert p.g_n[0] == pytest.approx(34000.0)
    lines, _ = budget_spending(p)
    assert lines["rent"][0] == pytest.approx(9000.0)


@pytest.mark.parametrize("given", [None, 0])
def test_net_spending_unset_or_zero_spends_the_budget_and_is_not_stored(given):
    """The UI always passes netSpending (0 by default); the default is not saved with the case."""
    p = _plan(state="", lines=[_core(50000.0), _rent(18000.0)], taxable=(800, 400))
    p.setSocialSecurity([2800, 2200], [70, 70])
    opts = dict(EXACT) if given is None else {**EXACT, "netSpending": given}
    p.solve("maxBequest", options=opts)
    assert p.g_n[0] == pytest.approx(68000.0)
    assert p.solverOptions.get("netSpending", None) == given


def test_budget_profile_needs_budget_lines():
    p = _plan(state="", lines=None)
    p.setSocialSecurity([2800, 2200], [70, 70])
    with pytest.raises(ValueError, match="Budget sheet"):
        p.solve("maxSpending", options={**EXACT, "bequest": 0})


def test_budget_lines_with_another_profile_warn_and_are_not_used():
    log = io.StringIO()
    p = _plan(state="", lines=[_core(50000.0), _rent(18000.0)], profile="flat", log=log)
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.solve("maxSpending", options={**EXACT, "bequest": 0})
    assert "Budget sheet is not used" in log.getvalue()
    assert p.budget is None and np.all(p.st_ptd_share_n == 0)


def test_budget_set_before_the_sheet_is_read_at_solve():
    """clone() sets the profile first and copies the household tables afterwards."""
    p = _plan(state="", lines=None, taxable=(800, 400))
    assert np.all(p.xi_n == 1.0)
    p.houseLists["Budget"] = _df(_core(50000.0), _line(name="care", type="care", year=THISYEAR + 10,
                                                       amount=20000.0))
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.solve("maxBequest", options=dict(EXACT))
    assert p.xi_n[10] == pytest.approx(70000.0 / 50000.0 * (1.0 if p.n_d > 10 else 0.6), rel=1e-9)


@pytest.mark.toml
def test_clone_with_a_new_lifespan_moves_the_survivor_step():
    from owlplanner import clone

    case = Path(__file__).resolve().parents[2] / "examples" / "Case_jack+jill.toml"
    p = owl.readConfig(str(case), verbose=False, logstreams=[io.StringIO()])
    p.setSpendingProfile("budget", 60)
    p.houseLists["Budget"] = _df(_core(80000.0), _rent(20000.0))
    p._evaluateBudget()
    c = clone(p, expectancy=[int(e) - 3 for e in p.expectancy], verbose=False)
    c._evaluateBudget()
    for q in (p, c):
        core = q.budget.by_type("core")
        assert core[q.n_d - 1] == pytest.approx(80000.0) and core[q.n_d] == pytest.approx(48000.0)
    assert c.n_d != p.n_d


def test_legacy_housing_sheet_is_read_as_budget(tmp_path):
    from openpyxl import load_workbook

    log = io.StringIO()
    p1 = _plan(state="", lines=[_core(50000.0), _rent(18000.0)])
    fname = str(tmp_path / "HFP_old.xlsx")
    p1.saveHFP(fname, overwrite=True)
    wb = load_workbook(fname)
    ws = wb["Budget"]
    ws.title = "Housing"
    ws.delete_cols([c.value for c in ws[1]].index("survivor") + 1)
    wb.save(fname)
    p2 = _plan(state="", lines=None, log=log)
    p2.readHFP(fname)
    assert list(p2.houseLists["Budget"]["name"]) == ["core", "rent"]
    assert p2.houseLists["Budget"]["survivor"].isna().all()
    assert "Housing sheet as the Budget sheet" in log.getvalue()


def test_hfp_round_trip_keeps_blank_and_set_survivor(tmp_path):
    lines = [_core(50000.0), _rent(18000.0) | {"survivor": 80.0, "rate": 0.5},
             _line(name="old", type="car", active=False)]
    p1 = _plan(state="", lines=lines)
    fname = str(tmp_path / "HFP_budget.xlsx")
    p1.saveHFP(fname, overwrite=True)
    p2 = _plan(state="", lines=None)
    p2.readHFP(fname)
    b = p2.houseLists["Budget"]
    assert list(b["active"]) == [True, True, False]
    assert np.isnan(b["survivor"].iloc[0]) and b["survivor"].iloc[1] == 80.0
    assert b["rate"].iloc[1] == pytest.approx(0.5)  # not read as 50%


def test_unknown_budget_type_is_named_in_a_warning(tmp_path):
    from openpyxl import load_workbook

    log = io.StringIO()
    p1 = _plan(state="", lines=[_core(50000.0)])
    fname = str(tmp_path / "HFP_typo.xlsx")
    p1.saveHFP(fname, overwrite=True)
    wb = load_workbook(fname)
    wb["Budget"].append([True, "boat", "yacht", int(THISYEAR), 0, 9000, 0, None])
    wb.save(fname)
    p2 = _plan(state="", lines=None, log=log)
    p2.readHFP(fname)
    assert list(p2.houseLists["Budget"]["name"]) == ["core"]
    assert "unknown type are left out: ['boat']" in log.getvalue()


def test_lines_left_out_warn_at_solve():
    log = io.StringIO()
    p = _plan(state="", lines=[_core(50000.0), _line(name="typo", type="car", end=10)], log=log)
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.solve("maxBequest", options=dict(EXACT))
    assert "Budget line 'typo' ends in 10" in log.getvalue()


# ---------------------------------------------------------------------------
# NJ property tax deduction, read from the budget
# ---------------------------------------------------------------------------


def test_nj_property_tax_deduction_owner():
    """A homeowner paying $20,000 property tax deducts the $15,000 cap."""
    p = _plan(state="NJ", lines=[_core(40000.0), _pt()], taxable=(400, 200))
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.setPension([4000, 3000], [62, 62], [False, False])
    p.solve("maxBequest", options=_bequest_opts())
    assert p.caseStatus == "solved"
    assert p.st_ptd_n[0] == pytest.approx(15000.0)
    assert p.st_pt_n[0] == pytest.approx(15000.0, abs=1.0)
    assert p.st_ti_n[0] < p.st_agi_n[0]


def test_nj_property_tax_deduction_tenant():
    """A tenant paying $30,000 rent deducts 18% of rent ($5,400)."""
    p = _plan(state="NJ", lines=[_core(40000.0), _rent(30000.0)], taxable=(400, 200))
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.setPension([4000, 3000], [62, 62], [False, False])
    p.solve("maxBequest", options=_bequest_opts())
    assert p.caseStatus == "solved"
    assert p.st_ptd_share_n[0] == pytest.approx(5400.0 / 70000.0)
    assert p.st_pt_n[0] == pytest.approx(5400.0, abs=1.0)


def test_deduction_follows_spending_under_max_spending():
    """The deductible share of net spending scales with it: at most share x g_n, at most the cap."""
    p = _plan(state="NJ", lines=[_core(40000.0), _pt(4000.0)], taxable=(900, 500))
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.setPension([4000, 3000], [62, 62], [False, False])
    p.solve("maxSpending", options={**_bequest_opts(), "bequest": 0})
    assert p.caseStatus == "solved"
    bound = np.minimum(p.st_ptd_n, p.st_ptd_share_n * p.g_n)
    assert np.all(p.st_pt_n <= bound + 1.0)
    assert p.st_T_n[0] > 1.0 and p.g_n[0] > 44000.0  # taxed, and spending above the budget's first year
    # Between the line's own amount and the cap: the deduction scaled with g_n, and the share binds.
    assert 4000.0 < bound[0] < 15000.0
    assert p.st_pt_n[0] == pytest.approx(bound[0], abs=1.0)


def test_ny_has_no_property_tax_deduction():
    p = _plan(state="NY", lines=[_core(40000.0), _pt()], taxable=(400, 200))
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.setPension([4000, 3000], [62, 62], [False, False])
    p.solve("maxBequest", options=_bequest_opts())
    assert p.caseStatus == "solved"
    assert np.all(p.st_ptd_n == 0) and np.all(p.st_pt_n == 0)


def test_move_to_new_jersey_deducts_only_there():
    p = _plan(state="NJ", lines=[_core(40000.0), _pt()], taxable=(400, 200))
    p.setStateTax("NJ", [(THISYEAR + 3, "FL")])
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.setPension([4000, 3000], [62, 62], [False, False])
    p.solve("maxBequest", options=_bequest_opts())
    assert p.caseStatus == "solved"
    assert np.all(p.st_ptd_n[:3] == 15000.0) and np.all(p.st_ptd_n[3:] == 0)


def _pension_only(lines):
    p = owl.Plan(["Joe", "Jane"], ["1962-06-15", "1963-12-15"], [85, 85], "NJ ptd", verbose=False)
    p.setSpendingProfile("budget", 60)
    p.setAccountBalances(taxable=[0, 0], taxDeferred=[0, 0], taxFree=[0, 0])
    p.setAllocationRatios("individual", generic=np.array([[[60, 40, 0, 0], [60, 40, 0, 0]]] * 2))
    p.setRates("conservative")
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.setPension([5000, 4000], [62, 62], [False, False])
    p.houseLists["Budget"] = _df(*lines)
    p.setStateTax("NJ")
    p.solve("maxBequest", options=_bequest_opts())
    assert p.caseStatus == "solved"
    return p


def test_deduction_does_not_change_the_exclusion_tier():
    """Line 41 comes after line 39: the tier is set on income before the property tax deduction."""
    p = _pension_only([_core(40000.0), _pt()])
    assert p.st_pt_n[0] == pytest.approx(15000.0, abs=1.0)
    # Line 27 (st_agi_n) is about $109k, in the 50% tier; the deduction would take it to about $94k,
    # in the 100% tier, if it moved the tier. The exclusion claimed must be the statute's, set on line 27.
    share, implied = p._tiered_exclusion_implied()
    assert 100000 < p.st_agi_n[0] < 115000 and share[0] == 0.5
    assert p.st_rx_n[0] == pytest.approx(implied[0], abs=2.0)
    statutory_ti = max(0.0, np.round(p.st_agi_n[0]) - p.st_rx_n[0] - p.st_sigmaBar_n[0])
    assert p.st_ti_n[0] == pytest.approx(statutory_ti - 15000.0, abs=2.0)


def test_local_search_runs_with_the_deduction():
    p = _plan(state="NJ", lines=[_core(40000.0), _pt()], taxable=(400, 200))
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.solve(
        "maxSpending",
        options={"breakpointMethod": "local-search", "localSearchTime": 1, "withMedicare": "None", "bequest": 0},
    )
    assert p.caseStatus == "solved"
    assert p.st_pt_n[0] <= min(p.st_ptd_n[0], p.st_ptd_share_n[0] * p.g_n[0]) + 1.0


def test_only_new_jersey_has_the_property_tax_deduction():
    for s in tax_state.valid_states():
        for filing in (0, 1):
            entry = tax_state.get_state_entry(s, filing)
            assert ("property_tax_deduction" in entry) == (s == "NJ")


def test_replayed_rows_match_a_fresh_build_with_nj_budget():
    """The loop-invariant builders are replayed after the first iteration (upstream #151)."""
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "fixed_rows", Path(__file__).parents[1] / "solver" / "test_fixed_rows.py"
    )
    fr = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(fr)

    p = _plan(state="NJ", lines=[_core(40000.0), _pt()], taxable=(400, 200))
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.setPension([4000, 3000], [62, 62], [False, False])
    p.solve("maxBequest", options=_bequest_opts())
    assert p.caseStatus == "solved"
    assert p.st_pt_n.sum() > 0
    assert p._fixedRows
    p._buildConstraints(p.objective, p.solverOptions)
    replayed = fr._lp_arrays(p)
    p._fixedRows = {}
    p._buildConstraints(p.objective, p.solverOptions)
    fr._assert_same_lp(fr._lp_arrays(p), replayed)


def test_process_before_any_solve():
    """The UI computes the fixed-asset bequest before the first solve (owlbridge.getFixedAssetsBequestValue)."""
    p = _plan(state="NJ", lines=[_core(40000.0), _pt()])
    fa = pd.DataFrame([{"active": True, "name": "home", "type": "residence", "year": THISYEAR, "basis": 500000.0,
                        "value": 800000.0, "rate": 0.0, "yod": 0, "commission": 5.0}])
    p.houseLists["Fixed Assets"] = conditionDebtsAndFixedAssetsDF(fa, "Fixed Assets")
    p.processDebtsAndFixedAssets()
    assert p.fixed_assets_bequest_value > 0
    assert np.all(p.st_ptd_n == 0)  # the state rule is only known once solve() reads the state


def test_metrics_and_explanation_report_housing_and_the_deduction():
    from owlplanner.assistant.explain import build_explanation
    from owlplanner.assistant.explain_schema import PlanExplanation

    p = _plan(state="NJ", lines=[_core(40000.0), _pt()], taxable=(400, 200))
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.setPension([4000, 3000], [62, 62], [False, False])
    p.solve("maxBequest", options={**_bequest_opts(), "withDuals": True})
    m = plan_metrics(p)
    assert m["housing_costs_today"] == pytest.approx(20000.0 * p.N_n, rel=1e-6)
    ex = build_explanation(p)
    PlanExplanation.model_validate(ex)
    assert ex["this_year"]["state_tax"]["property_tax_deduction"] == pytest.approx(15000.0, abs=1.0)
    by_year = ex["state_tax_brackets"]["by_year"]
    assert any(r.get("property_tax_deduction_today", 0) > 0 for r in by_year)


# ---------------------------------------------------------------------------
# Essential lines: a floor paid at its amount, the other lines scaled (fork, after #175)
# ---------------------------------------------------------------------------


def _ess(line):
    return line | {"essential": True}


def _travel(amount=20000.0, **kw):
    return _line(name="travel", type="travel", amount=amount, **kw)


def _df_ess(*lines):
    return conditionDebtsAndFixedAssetsDF(pd.DataFrame(list(lines), columns=COLS + ["essential"]), "Budget")


def _plan_ess(lines, taxable=(800, 400), expectancy=(89, 92)):
    p = owl.Plan(["Joe", "Jane"], ["1964-03-15", "1965-09-15"], list(expectancy), "essential", verbose=False,
                 logstreams=[io.StringIO()])
    p.setSpendingProfile("budget", 60)
    p.setAccountBalances(taxable=list(taxable), taxDeferred=[0, 0], taxFree=[0, 0])
    p.setAllocationRatios("individual", generic=np.array([[[60, 40, 0, 0], [60, 40, 0, 0]]] * 2))
    p.setRates("conservative")
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.houseLists["Budget"] = _df_ess(*lines)
    return p


class TestEssentialFlag:
    @pytest.mark.parametrize("val,expected", [(True, True), ("yes", True), (1, True), ("True", True),
                                              (False, False), ("no", False), (0, False), (np.nan, False),
                                              (None, False), ("", False)])
    def test_blank_means_no(self, val, expected):
        assert budget.is_essential({"essential": val}) is expected

    def test_column_is_optional_and_blank_is_false(self):
        df = _df(_core(50000.0))
        assert list(df["essential"]) == [False]
        df = _df_ess(_core(50000.0) | {"essential": np.nan}, _ess(_rent(18000.0)))
        assert list(df["essential"]) == [False, True]

    def test_evaluate_splits_essential_and_discretionary(self):
        b = budget.evaluate(_df_ess(_ess(_core(50000.0)), _ess(_rent(18000.0)), _travel()), 6, 3, 60, THISYEAR)
        assert b.has_essentials and b.essential == (True, True, False)
        assert list(b.essential_n) == [68000.0] * 3 + [48000.0] * 3  # rent kept in full by a survivor
        assert list(b.discretionary_n) == [20000.0] * 3 + [12000.0] * 3
        assert not budget.evaluate(_df(_core(50000.0)), 6, 3, 60, THISYEAR).has_essentials

    def test_hfp_round_trip_keeps_the_flag(self, tmp_path):
        p1 = _plan_ess([_ess(_core(50000.0)), _travel()])
        fname = str(tmp_path / "HFP_essential.xlsx")
        p1.saveHFP(fname, overwrite=True)
        p2 = _plan(state="", lines=None)
        p2.readHFP(fname)
        assert list(p2.houseLists["Budget"]["essential"]) == [True, False]


def test_max_spending_pays_essentials_and_scales_the_rest():
    p = _plan_ess([_ess(_core(50000.0)), _ess(_rent(18000.0)), _travel(end=THISYEAR + 9)])
    p.solve("maxSpending", options={**EXACT, "bequest": 0})
    assert p.caseStatus == "solved"
    k = p.discretionary_scale
    assert k > 1  # the couple can afford more than the budgeted travel
    real = p.g_n / p.gamma_n[:-1]
    assert real == pytest.approx(p.budget.essential_n + k * p.budget.discretionary_n, rel=1e-6)
    lines, housing = budget_spending(p)
    assert sum(lines.values()) == pytest.approx(p.g_n, rel=1e-6)  # after travel ends: g_n = E_n to solver tolerance
    assert lines["rent"] == pytest.approx(18000.0 * p.gamma_n[:-1], rel=1e-6)  # not scaled
    assert lines["travel"][0] == pytest.approx(20000.0 * k, rel=1e-6)
    assert housing == pytest.approx(lines["rent"])
    assert "Discretionary spending (share of budget)" in p.summaryDic()
    assert plan_metrics(p)["discretionary_scale"] == pytest.approx(k)


def test_without_essential_lines_the_plan_is_unchanged():
    lines = [_core(50000.0), _rent(18000.0), _travel(end=THISYEAR + 9)]
    p_a = _plan(state="", lines=lines, taxable=(800, 400))
    p_b = _plan_ess(lines)
    for p in (p_a, p_b):
        p.setSocialSecurity([2800, 2200], [70, 70])
        p.solve("maxSpending", options={**EXACT, "bequest": 0})
    assert p_b.g_n == pytest.approx(p_a.g_n, rel=1e-9)
    assert p_b.discretionary_scale is None and "discretionary_scale" not in plan_metrics(p_b)


def test_max_bequest_at_the_budget_is_the_same_plan():
    """At the budgeted level (scale 1) the flag changes nothing: the spending path is the budget."""
    plain = [_core(50000.0), _rent(18000.0), _travel(end=THISYEAR + 9)]
    p_a = _plan_ess(plain)
    p_b = _plan_ess([_ess(plain[0]), _ess(plain[1]), plain[2]])
    for p in (p_a, p_b):
        p.solve("maxBequest", options=dict(EXACT))
    assert p_b.discretionary_scale == pytest.approx(1.0)
    assert p_b.g_n == pytest.approx(p_a.g_n, rel=1e-9)
    assert _bequest(p_b) == pytest.approx(_bequest(p_a), rel=1e-9)


def test_max_bequest_on_essentials_only():
    """netSpending at the essential level: no discretionary spending, the most the heirs can get."""
    p = _plan_ess([_ess(_core(50000.0)), _ess(_rent(18000.0)), _travel()])
    p.solve("maxBequest", options={**EXACT, "netSpending": 68})
    assert p.discretionary_scale == pytest.approx(0.0, abs=1e-9)
    assert p.g_n / p.gamma_n[:-1] == pytest.approx(p.budget.essential_n, rel=1e-6)


def test_max_bequest_below_the_essentials_is_refused():
    p = _plan_ess([_ess(_core(50000.0)), _travel()])
    with pytest.raises(ValueError, match="below the budget's essential lines"):
        p.solve("maxBequest", options={**EXACT, "netSpending": 40})


def test_every_line_essential_is_refused_under_max_spending():
    p = _plan_ess([_ess(_core(50000.0)), _ess(_rent(18000.0))])
    with pytest.raises(ValueError, match="need a discretionary line"):
        p.solve("maxSpending", options={**EXACT, "bequest": 0})


def test_essentials_beyond_the_means_are_infeasible():
    """Whole-budget scaling would cut rent with travel; a floor the savings cannot pay is reported."""
    lines = [_core(60000.0), _rent(30000.0), _travel()]
    p_a = _plan_ess(lines, taxable=(150, 100))
    p_a.solve("maxSpending", options={**EXACT, "bequest": 0})
    assert p_a.caseStatus == "solved" and p_a.g_n[0] < 110000.0  # everything scaled down, rent too
    p_b = _plan_ess([_ess(lines[0]), _ess(lines[1]), lines[2]], taxable=(150, 100))
    p_b.solve("maxSpending", options={**EXACT, "bequest": 0})
    assert p_b.caseStatus != "solved"


def test_essential_rows_replay_across_iterations():
    """The profile rows are cached across loop iterations; the loop's plan keeps the floor."""
    p = _plan_ess([_ess(_core(50000.0)), _ess(_rent(18000.0)), _travel()])
    p.solve("maxSpending", options={"bequest": 0})  # default loop: Medicare and SS taxability iterate
    assert p.caseStatus == "solved"
    real = p.g_n / p.gamma_n[:-1]
    k = p.discretionary_scale
    assert real == pytest.approx(p.budget.essential_n + k * p.budget.discretionary_n, rel=1e-6)


def test_spending_slack_applies_to_the_discretionary_part():
    p = _plan_ess([_ess(_core(50000.0)), _travel()])
    p.solve("maxSpending", options={**EXACT, "bequest": 0, "spendingSlack": 10})
    real = p.g_n / p.gamma_n[:-1]
    disc = real - p.budget.essential_n
    ratio = disc / p.budget.discretionary_n / (disc[0] / p.budget.discretionary_n[0])
    assert np.all(ratio >= 0.9 - 1e-6) and np.all(ratio <= 1.1 + 1e-6)
    assert np.all(real >= p.budget.essential_n - 1e-3)
