"""
Phase 5 step 6: sweep work (and other) variants of one case, deciding claiming ages by the
two-stage method of `fork-notes/plan-review-2026-10-22.md` §3.2 (decision R1, 2026-10-22).

For each variant:

(a) reference search: `withSSAges` under the exact LP (`withMedicare="none"`,
    `withSSTaxability=0.85`), with `withACA="optimize"` added when the case has a benchmark
    premium (the same rule as the local-search preset), so the ACA cliff is priced;
(b) loop search: `withSSAges` with every breakpoint family on the loop (skip with --no-loop);
(c) decision at fixed ages. For each searched person in turn, with the others held at the current
    best: the whole years 62-70 (each mapped to the first open month at or after it), the picks of
    (a) and (b) and the current best. The best score becomes that person's age; passes repeat
    while someone moves (--passes, default 2).

Scoring in (c), --decider (2026-10-26, the user's call after `ws_bb_check.py`, revising R1):

- "bb" (default): branch-and-bound on the case's own model (`breakpointMethod="branch-and-bound"`,
  every family exact) at --gap (default 1e-5; without a gap option, `withMedicare="optimize"`
  loosens Owl's gap to 3e-3, about $340/yr on chris+pat) and --max-time per MIP (Owl's 900 s). A
  plan is *proved* when every HiGHS run ended "Optimal" and none stopped on time or at the node
  limit; its *tolerance* is the accepted solve's reported gap x score, the distance to the model's
  optimum at those ages that HiGHS's bound allows. The reported gap can exceed --gap (1.0-2.0e-5
  on 12 of 84 chris+pat plans at --gap 1e-5): HiGHS stops within --gap, but the branch-and-bound
  loop's sticky binaries (upstream `6250a5b0`) may keep the previous iteration's plan when it is
  within the loop's relTol (5e-5 by default) of HiGHS's, and re-express the bound against it
  (`highs_gap_probe.py`). So --gap alone bounded the plan to --gap + relTol; the script passes
  relTol = --gap (2026-10-26, the user's call), which bounds it to 2 x --gap (on 15 chris+pat
  plans it changed nothing: `reltol_probe_2026-10-26.txt`). Upstream fixed this as #233 (`4aa4887`,
  merged 2026-10-28): a tie is refused when it would take a plan HiGHS proved within --gap outside
  it, so --gap now bounds the plan by itself; relTol = --gap is still passed (it only narrows
  which ties are kept). An unproved plan is also scored by
  local search and the better of the two kept, labelled by the method that produced it. On
  chris+pat branch-and-bound at 1e-5 beat local search by $15-74/yr on all 6 plans checked, in
  94-288 s each (`ws_bb_check_2026-10-26.txt`); HiGHS branch-and-bound can be far slower on other
  cases.
- "ls": local search at fixed ages, its time caps raised (--ls-time, --ls-step-time) so its node
  limits decide; scores carry search noise.

Resolution. With proved plans, a difference is resolved when it is larger than the two plans'
tolerances together; the report also names sweeps whose best and second-best ages are that close,
and a dip between proved plans is the model's own (on chris+pat, part time, Pat 64y0m). Where a
score came from local search, the noise is estimated as before: assuming the true value is
unimodal in each claiming age, an interior dip (a score below both neighbours) is a lower bound on
the noise, and the deepest one is the variant's spread; it can also be the model's own (the base
variant's dip at Pat 64y0m was local search's, the part-time one the model's). The per-plan scores
in --json show which costs move at a dip.

Two checks on the decided plans (2026-10-26, after upstream #234):

- Round-trip netting. Owl's worksheet nets a taxable sale against the surplus deposit that rebuys
  it in years with no federally taxed gains (`Plan._netSurplusRoundTrip`), while the gains, the
  state tax and the cost basis keep the gross flows. Since upstream fixed #234 (`3931c8a`, merged
  2026-10-28), only a round trip that realizes no gain is netted, so the hidden gain recorded
  below should be 0; it stays recorded as a check. Every scored plan records what was netted
  per year, in today's $: the amount, and the gain it hides (amount x gain coefficient x equity
  share, the part of `Q_n` the displayed flows no longer explain), with that year's state tax.
  NY taxes the gain, so a netted NY year carries state tax on a sale the worksheet does not show.
- Deposit split. The surplus goes to the spouses' taxable accounts in a fixed split (`eta`,
  `spousal_surplus_deposit_fraction`), a modelling constraint, not a household fact; a forced
  split can make the plan sell in one account to deposit in the other (#234). For couples, each
  variant's decided ages are scored again at every --eta value (default 0 and 1) with the same
  decider, and the report gives the change against the case's own split. The ages are not
  re-decided: a difference beyond the tolerances says the split matters, not which ages it picks.

Variants are generic case and HFP overrides, so the script also serves the locality and rent/buy
comparisons (Phase 6 folded in, decision R5). A JSON file describes them:

    {
      "set": ["aca_settings.slcsp_annual=20", "solver_options.withSSAges=optimize"],
      "variants": [
        {"name": "base", "set": ["aca_settings.aca_start_year=2029"]},
        {"name": "stops-2027",
         "set": ["aca_settings.aca_start_year=2028"],
         "hfp": {"Pat": [{"years": [2028, 2028], "anticipated wages": 0, "401k ctrb": 0}]}}
      ]
    }

- "set": `owlcli --set` specs (KEY.PATH=VALUE, the same parser), the top-level list applied to
  every variant first, then the variant's own. A PIA per variant goes here
  (`fixed_income.social_security_pia_amounts=[...]`, decision D4).
- "hfp": per person (sheet name), edits of that person's sheet: "years" [first, last] and any
  recognized column with its value in dollars, as typed in the workbook. Applied with openpyxl to
  a copy of the variant's workbook, so every other cell keeps its value and format (the fork reads
  Debts/Fixed Assets rates by their percent format). Wages follow upstream #221: the wages in
  AGI, pay less pre-tax deferrals.
- The first variant is the reference for the differences.

Each variant's workbook and case file are written to --out-dir, so any run can be repeated with
`owlcli run`. Real numbers stay in `otherFiles/` (the default --out-dir is `otherFiles/sweep`).

Scores should not depend on machine load or --jobs: branch-and-bound's are load-free unless a
MIP stops on --max-time (reported per plan); local search runs with its time caps raised and the
report says where its default caps (300 s per solve, 60 s per restricted solve) would have bound.

Plans start today: compare variants run the same day. Run from the repository root:

    .venv/bin/python fork-notes/phase5/work_sweep.py CASE.toml VARIANTS.json [--out-dir DIR]
        [--jobs N] [--json FILE]
"""

import argparse
import copy
import io
import json
import os
import sys
import time
from concurrent.futures import ProcessPoolExecutor
from datetime import date
import multiprocessing as mp

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "src"))

import numpy as np  # noqa: E402
import openpyxl  # noqa: E402

import owlplanner as owl  # noqa: E402
from owlplanner import hfp_io, localsearch  # noqa: E402
from owlplanner.cli.set_override import apply_overrides  # noqa: E402
from owlplanner.config import load_toml, save_toml  # noqa: E402
from owlplanner.config.plan_bridge import config_to_plan  # noqa: E402
from owlplanner.export import plan_metrics  # noqa: E402
from owlplanner.plan import TIME_LIMIT, Plan  # noqa: E402

EXACT = {"withMedicare": "none", "withSSTaxability": 0.85}
FAMILIES = ("withSSTaxability", "withLTCG", "withNIIT", "withACA", "withSeniorBonus", "withMedicare")
PRESET_KEYS = ("breakpointMethod", "mipStrategy", "localSearchTime", "localSearchStepTime",
               "localSearchRadius", "localSearchStepNodes")
DECIDER = {"bb": "branch-and-bound", "ls": "local search"}
METRICS = ("federal_income_tax_today", "ltcg_tax_today", "niit_today", "state_tax_today", "payroll_tax_today",
           "medicare_today", "aca_today", "ss_withheld_today", "ss_income_today", "final_bequest_today")


def ym(age):
    m = int(round(age * 12))
    return f"{m // 12}y{m % 12}m"


def months(ages):
    return tuple(int(round(a * 12)) for a in ages)


# ----- building a variant -------------------------------------------------------------------------
def _hfp_path(diconf, dirname):
    name = diconf.get("household_financial_profile", {}).get("HFP_file_name", "None")
    if not name or name == "None":
        return None
    if os.path.exists(name):
        return os.path.abspath(name)
    path = os.path.join(dirname, name)
    if not os.path.exists(path):
        raise FileNotFoundError(f"HFP '{name}' not found.")
    return os.path.abspath(path)


def edit_hfp(src, dst, edits):
    """Copy workbook src to dst with the per-person edits applied cell by cell."""
    known = set(hfp_io.timeHorizonItems()) - {"year"}
    wb = openpyxl.load_workbook(src)
    for person, items in edits.items():
        if person not in wb.sheetnames:
            raise ValueError(f"HFP edit: no sheet '{person}' in {src}.")
        ws = wb[person]
        header = {c.value: c.column for c in ws[1] if c.value is not None}
        if "year" not in header:
            raise ValueError(f"HFP edit: sheet '{person}' has no 'year' column.")
        rows = {ws.cell(row=r, column=header["year"]).value: r for r in range(2, ws.max_row + 1)}
        for item in items:
            first, last = item["years"]
            cols = {k: v for k, v in item.items() if k != "years"}
            bad = [k for k in cols if k not in known]
            if bad:
                raise ValueError(f"HFP edit for {person}: unknown column(s) {bad}; known: {sorted(known)}.")
            for col in cols:
                if col not in header:
                    header[col] = max(header.values()) + 1
                    ws.cell(row=1, column=header[col], value=col)
            for year in range(int(first), int(last) + 1):
                if year not in rows:
                    rows[year] = ws.max_row + 1
                    ws.cell(row=rows[year], column=header["year"], value=year)
                for col, value in cols.items():
                    ws.cell(row=rows[year], column=header[col], value=value)
    wb.save(dst)


def build_variants(case_file, spec, out_dir):
    """Return [(name, diconf, dirname, overrides)] with each variant's workbook written."""
    base, dirname, _ = load_toml(case_file)
    os.makedirs(out_dir, exist_ok=True)
    out = []
    names = set()
    for v in spec["variants"]:
        name = v["name"]
        if name in names:
            raise ValueError(f"Duplicate variant name '{name}'.")
        names.add(name)
        sets = list(spec.get("set", [])) + list(v.get("set", []))
        d = apply_overrides(base, sets)
        stem = f"{d.get('case_name', 'case')}_{name}"
        hfp = _hfp_path(d, dirname)
        if v.get("hfp"):
            if hfp is None:
                raise ValueError(f"Variant '{name}' edits the HFP but the case has none.")
            dst = os.path.abspath(os.path.join(out_dir, f"HFP_{stem}.xlsx"))
            edit_hfp(hfp, dst, v["hfp"])
            hfp = dst
        if hfp is not None:
            d.setdefault("household_financial_profile", {})["HFP_file_name"] = hfp
        save_toml(d, os.path.join(out_dir, f"Case_{stem}.toml"))
        out.append((name, d, dirname, sets + ([f"hfp: {json.dumps(v['hfp'])}"] if v.get("hfp") else [])))
    return out


def make_plan(diconf, dirname, ages=None, eta=None):
    d = copy.deepcopy(diconf)
    if ages is not None:
        d.setdefault("fixed_income", {})["social_security_ages"] = [m / 12 for m in months(ages)]
    if eta is not None:
        d.setdefault("savings_assets", {})["spousal_surplus_deposit_fraction"] = float(eta)
    buf = io.StringIO()
    p = config_to_plan(d, dirname, verbose=True, logstreams=[buf])
    return p, buf


def base_options(p, **extra):
    opts = {k: v for k, v in dict(p.solverOptions or {}).items() if k not in PRESET_KEYS}
    opts.update(extra)
    return opts


# ----- stages (a) and (b): age searches -----------------------------------------------------------
def age_search(diconf, dirname, mode):
    p, _ = make_plan(diconf, dirname)
    sel = (p.solverOptions or {}).get("withSSAges", "fixed")
    if mode == "exact":
        extra = dict(EXACT, withACA="optimize") if p.slcsp_annual > 0 else dict(EXACT)
    else:
        extra = {f: "loop" for f in FAMILIES}
        if (p.solverOptions or {}).get("withMedicare") == "none":
            extra["withMedicare"] = "none"
    t0 = time.time()
    p.solve(p.objective, base_options(p, withSSAges=sel, **extra))
    dt = time.time() - t0
    ok = p.caseStatus == "solved"
    return {"ages": [float(a) for a in p.ssecAges], "value": p._objectiveValue(p.objective) if ok else None,
            "status": p.caseStatus, "seconds": dt}


# ----- stage (c): scoring at fixed ages -----------------------------------------------------------
_HIGHS_RUNS = []  # (status, gap) of every HiGHS run in this worker's current solve


def _record_highs_runs():
    """Wrap Plan._run_highs once per worker process (one plan solves at a time in a worker)."""
    if getattr(Plan._run_highs, "_recorded", False):
        return
    run_highs = Plan._run_highs

    def recorded(self, *a, **k):
        res = run_highs(self, *a, **k)
        _HIGHS_RUNS.append((str(res[3]), float(res[4])))
        return res

    recorded._recorded = True
    Plan._run_highs = recorded


def _record_netting():
    """Wrap Plan._netSurplusRoundTrip once per worker process: each call leaves on the plan what it
    netted, per year in today's $. On the plan, not in a module list, so a state the solve restores
    afterwards (local search falling back to the loop's plan) carries its own record."""
    if getattr(Plan._netSurplusRoundTrip, "_recorded", False):
        return
    net = Plan._netSurplusRoundTrip

    def recorded(self):
        w0 = self.w_ijn[:, 0, :].copy()
        net(self)
        rows = []
        for n in range(self.N_n):
            delta = w0[:, n] - self.w_ijn[:, 0, n]
            if delta.sum() <= 0.5:
                continue
            g = float(self.gamma_n[n])
            hidden = sum(float(delta[i]) * self._effective_cap_gain_coef(i, n) * float(self.alpha_ijkn[i, 0, 0, n])
                         for i in range(self.N_i))
            basis = self.gain_fraction_in is not None and any(
                delta[i] > 0.5 and not np.isnan(self.gain_fraction_in[i, n]) for i in range(self.N_i))
            rows.append([int(self.year_n[n]), float(delta.sum()) / g, hidden / g, float(self.st_T_n[n]) / g, basis])
        self._ws_netted = rows

    recorded._recorded = True
    Plan._netSurplusRoundTrip = recorded


def _netting(p):
    """The netting record of a solved plan: totals and the years, today's $."""
    rows = getattr(p, "_ws_netted", None) or []
    return {"netted": sum(r[1] for r in rows), "hidden_gain": sum(r[2] for r in rows),
            "years": [{"year": r[0], "netted": r[1], "hidden_gain": r[2], "state_tax": r[3], "basis": r[4]}
                      for r in rows]}


def _solved(p, t0):
    return {"status": p.caseStatus, "seconds": time.time() - t0, "method": getattr(p, "breakpointMethodUsed", None)}


def _score_ls(diconf, dirname, ages, cfg):
    _record_netting()
    p, buf = make_plan(diconf, dirname, ages, cfg.get("eta"))
    opts = base_options(p, withSSAges="fixed", breakpointMethod="local-search",
                        localSearchTime=cfg["ls_time"], localSearchStepTime=cfg["ls_step"])
    t0 = time.time()
    p.solve(p.objective, opts)
    out = _solved(p, t0)
    out["used"] = "local search"
    if p.caseStatus != "solved":
        return out
    log = getattr(p, "localSearchLog", None) or []
    steps = [s["seconds"] for it in log for s in it.get("steps", [])]
    out.update({
        "value": p._objectiveValue(p.objective),
        "search_seconds": sum(it.get("time", 0.0) for it in log),
        "max_step_seconds": max(steps) if steps else 0.0,
        "budget_spent": "time budget spent" in buf.getvalue(),
        "metrics": {k: v for k, v in plan_metrics(p).items() if k in METRICS},
        "netting": _netting(p),
    })
    out["default_caps_bind"] = (out["search_seconds"] > localsearch.TOTAL_TIME
                                or out["max_step_seconds"] > localsearch.STEP_TIME)
    return out


def _score_bb(diconf, dirname, ages, cfg):
    _record_highs_runs()
    _record_netting()
    _HIGHS_RUNS.clear()
    p, _ = make_plan(diconf, dirname, ages, cfg.get("eta"))
    opts = base_options(p, withSSAges="fixed", breakpointMethod="branch-and-bound",
                        gap=cfg["gap"], relTol=cfg["gap"], maxTime=cfg["max_time"])
    t0 = time.time()
    p.solve(p.objective, opts)
    out = _solved(p, t0)
    runs = list(_HIGHS_RUNS)
    hits = getattr(p, "solverNodeLimitHits", None)
    out.update({
        "used": "branch-and-bound",
        "highs_runs": len(runs),
        "highs_statuses": sorted({st for st, _ in runs}),
        "time_limit_runs": sum("time limit" in st.lower() for st, _ in runs),
        "node_limit_runs": int(hits[0]) if hits else 0,
    })
    if p.caseStatus != "solved":
        out["proved"] = False
        return out
    out.update({
        "value": p._objectiveValue(p.objective), "gap": float(p.solverGap),
        "residual": p._fixedPointResidualTotal(),
        "sticky_tol": getattr(p, "_sticky_tol", None),
        "metrics": {k: v for k, v in plan_metrics(p).items() if k in METRICS},
        "netting": _netting(p),
    })
    return _prove(out)


def _prove(c):
    """Set proved / tolerance on a branch-and-bound cell from its recorded runs (also on re-report)."""
    ok = (c.get("status") == "solved" and c.get("highs_runs", 0) > 0
          and all(st == "Optimal" for st in c.get("highs_statuses", []))
          and not c.get("time_limit_runs") and not c.get("node_limit_runs")
          and c.get("gap") is not None and c["gap"] >= 0)
    c["proved"] = ok
    c["tolerance"] = c["gap"] * abs(c["value"]) if ok else None
    return c


def score(args):
    """Worker: one plan at fixed ages. Picklable arguments, a plain dict back."""
    diconf, dirname, ages, cfg = args
    if cfg["decider"] == "ls":
        out = _score_ls(diconf, dirname, ages, cfg)
    else:
        out = _score_bb(diconf, dirname, ages, cfg)
        if not out.get("proved"):
            ls = _score_ls(diconf, dirname, ages, cfg)
            out["fallback"] = {k: v for k, v in ls.items() if k != "metrics"}
            if ls.get("value") is not None and (out.get("value") is None or ls["value"] > out["value"]):
                keep = {k: out[k] for k in ("highs_runs", "highs_statuses", "time_limit_runs", "node_limit_runs",
                                            "gap", "fallback") if k in out}
                out = dict(ls, bb=dict(keep, value=out.get("value"), status=out["status"]), proved=False,
                           tolerance=None)
    out["ages"] = list(months(ages))
    return out


def whole_year_candidates(p, i):
    """First open claiming month at or after each whole year 62..70 for person i."""
    open_m = [int(round(a * 12)) for a in p._ssaOpenAges(i)]
    picks = []
    for y in range(62, 71):
        later = [m for m in open_m if m >= 12 * y]
        if later and later[0] not in picks:
            picks.append(later[0])
    return picks


def dips(points):
    """Interior dips of [(months, value)] sorted by age: [(months, depth)]."""
    pts = sorted(points)
    out = []
    for (_, a), (m, b), (_, c) in zip(pts, pts[1:], pts[2:], strict=False):
        if b < a and b < c:
            out.append((m, min(a, c) - b))
    return out


def decide(diconf, dirname, ref, loop, pool, cfg, passes, log):
    p, _ = make_plan(diconf, dirname)
    p._ssa_optimize_set = p._parseSSAgesOption(p.solverOptions or {})
    free = [i for i in sorted(p._ssa_optimize_set) if not p._ssaAgeIsFixed(i)]
    best = list(months(ref["ages"] if ref else p.ssecAges))
    cache = {}

    def run(trials):
        todo = [t for t in dict.fromkeys(trials) if t not in cache]
        jobs = [(diconf, dirname, [m / 12 for m in t], cfg) for t in todo]
        for t, r in zip(todo, pool.map(score, jobs), strict=True):
            cache[t] = r

    def val(t):
        return cache[t].get("value")

    run([tuple(best)])
    sweeps = []
    seen = set()
    for k in range(passes if free else 0):
        moved = False
        for i in free:
            key = (i, tuple(best[:i] + best[i + 1:]))  # the others' ages: same key, same sweep
            if key in seen:
                continue
            seen.add(key)
            opts = set(whole_year_candidates(p, i)) | {best[i]}
            for pick in (ref, loop):
                if pick and pick["status"] == "solved":
                    opts.add(months(pick["ages"])[i])
            trials = []
            for m in sorted(opts):
                t = list(best)
                t[i] = m
                trials.append(tuple(t))
            t0 = time.time()
            run(trials)
            pts = [(t[i], val(t)) for t in trials if val(t) is not None]
            winner = max(trials, key=lambda t: -float("inf") if val(t) is None else val(t))
            sweeps.append({"pass": k, "person": p.inames[i], "index": i, "held": list(best), "points": pts,
                           "dips": dips(pts)})
            log(f"    pass {k}, {p.inames[i]} swept ({len(trials)} ages, {time.time() - t0:.0f}s): "
                f"best {ym(winner[i] / 12)} at {val(winner):,.2f}")
            vb = val(tuple(best))
            if winner[i] != best[i] and val(winner) is not None and (vb is None or val(winner) > vb):
                best[i] = winner[i]
                moved = True
        if not moved:
            log(f"    pass {k}: no age moved")
            break
    return tuple(best), cache, sweeps, [p.inames[i] for i in free]


# ----- report -------------------------------------------------------------------------------------
def run_variant(name, diconf, dirname, args, pool, log):
    log(f"\n== {name}")
    ref = age_search(diconf, dirname, "exact")
    log(f"  (a) exact-LP search: {ref['status']}, ages {'/'.join(ym(a) for a in ref['ages'])} "
        f"({ref['seconds']:.0f}s)")
    loop = None
    if not args.no_loop:
        loop = age_search(diconf, dirname, "loop")
        log(f"  (b) loop search:     {loop['status']}, ages {'/'.join(ym(a) for a in loop['ages'])} "
            f"({loop['seconds']:.0f}s)")
    log(f"  (c) {DECIDER[args.decider]} at fixed ages:")
    cfg = {"decider": args.decider, "gap": args.gap, "max_time": args.max_time,
           "ls_time": args.ls_time, "ls_step": args.ls_step_time}
    best, cache, sweeps, free = decide(diconf, dirname, ref if ref["status"] == "solved" else None, loop, pool,
                                       cfg, args.passes, log)
    etas = eta_check(diconf, dirname, best, pool, cfg, args.eta, log)
    return summarize(name, free, best, cache, sweeps, ref, loop, etas)


def eta_check(diconf, dirname, best, pool, cfg, etas, log):
    """The decided ages scored again at each other deposit split: [[eta, cell]] (couples only)."""
    p, _ = make_plan(diconf, dirname)
    own = float(p.eta)
    others = [e for e in etas if abs(e - own) > 1e-9]
    if p.N_i != 2 or not others:
        return []
    t0 = time.time()
    ages = [m / 12 for m in best]
    cells = list(pool.map(score, [(diconf, dirname, ages, dict(cfg, eta=e)) for e in others]))
    log(f"  deposit split at the decided ages ({time.time() - t0:.0f}s): case's {own:.2f}; "
        + ", ".join(f"{e:.2f} -> {c.get('value', float('nan')):,.2f}" for e, c in zip(others, cells, strict=True)))
    return [[e, c] for e, c in zip(others, cells, strict=True)]


def summarize(name, free, best, cache, sweeps, ref, loop, etas=None):
    """A variant's row: decided plan, noise and tolerance measures, every scored plan."""
    uniq, seen = [], set()
    for s in sweeps:  # a sweep repeated with the same other ages adds nothing (older runs kept them)
        held = list(s["held"])
        held[s["index"]] = None
        key = (s["index"], tuple(held))
        if key not in seen:
            seen.add(key)
            uniq.append(s)
    sweeps = uniq
    # Noise: the dips matter only where a score is local search's (an unproved plan).
    ls_dips = [d for s in sweeps for m, d in s["dips"] if not _all_proved(cache, s, m)]
    spread = max(ls_dips, default=None)
    close = []
    for s in sweeps:
        ranked = sorted(((v, m) for m, v in s["points"]), reverse=True)
        if len(ranked) > 1:
            (v1, m1), (v2, m2) = ranked[0], ranked[1]
            t1, t2 = (_cell(cache, s, m).get("tolerance") for m in (m1, m2))
            if t1 is not None and t2 is not None and v1 - v2 <= t1 + t2:
                close.append(f"{s['person']} {ym(m1 / 12)} vs {ym(m2 / 12)} ({v1 - v2:,.2f} within "
                             f"{t1 + t2:,.2f})")

    def at(pick):
        if not pick or pick["status"] != "solved":
            return None
        return cache.get(months(pick["ages"]), {}).get("value")

    return {"name": name, "free": free, "best": list(best), "result": cache[tuple(best)], "spread": spread,
            "close": close, "ref": ref, "loop": loop, "ref_value": at(ref), "loop_value": at(loop), "sweeps": sweeps,
            "results": {"/".join(ym(m / 12) for m in k): v for k, v in cache.items()}, "_cache": cache,
            "eta": etas or []}


def reload(path):
    """Rows of a saved --json run, re-classified with the current proof rule (no solving)."""
    with open(path, encoding="utf-8") as f:
        saved = json.load(f)
    rows = []
    for r in saved:
        cache = {}
        for c in r["results"].values():
            if "highs_statuses" in c and c.get("used") == "branch-and-bound":
                _prove(c)
            cache[tuple(c["ages"])] = c
        etas = [[e, _prove(c) if c.get("used") == "branch-and-bound" and "highs_statuses" in c else c]
                for e, c in r.get("eta", [])]
        rows.append(summarize(r["name"], r["free"], r["best"], cache, r["sweeps"], r["ref"], r["loop"], etas))
    return rows


def _cell(cache, sweep, m):
    """The scored plan of a sweep at person's age m (months)."""
    i = sweep["index"]
    t = list(sweep["held"])
    t[i] = m
    return cache[tuple(t)]


def _all_proved(cache, sweep, m):
    """True when the dip at m and both its neighbours in the sweep are proved plans."""
    ms = sorted(x for x, _ in sweep["points"])
    k = ms.index(m)
    return all(_cell(cache, sweep, x).get("proved") for x in ms[k - 1:k + 2])


def _verdict(r, r0):
    """Resolution of r's difference from the reference row r0."""
    a, b = r["result"], r0["result"]
    d = a["value"] - b["value"]
    if a.get("tolerance") is not None and b.get("tolerance") is not None:
        return "resolved" if abs(d) > a["tolerance"] + b["tolerance"] else "unresolved"
    sp, sp0 = r["spread"], r0["spread"]
    if sp is None or sp0 is None:
        return "no spread"
    return "resolved" if abs(d) > max(sp, sp0) else "unresolved"


def report(rows, objective, decider, log):
    unit = "annual spending, today's $/yr" if objective == "maxSpending" else "after-tax bequest, today's $"
    log(f"\nDecided by {DECIDER[decider]}; objective = {unit}. Taxes and premiums: lifetime, today's $.")
    base = rows[0]["result"].get("value")
    hdr = (f"{'variant':<24}{'ages':>14}{'objective':>13}{'vs first':>11}{'tol':>7}{'spread':>8}  {'':<11}"
           f"{'(a) cost':>9}{'(b) cost':>9}  by")
    log(hdr)
    for r in rows:
        res = r["result"]
        v = res.get("value")
        if v is None:
            log(f"{r['name']:<24}{'':>14}{res['status']:>13}")
            continue
        sp = r["spread"]
        tol = res.get("tolerance")
        d = v - base if base is not None else None
        verdict = "" if r is rows[0] else _verdict(r, rows[0])
        costs = [f"{v - x:>9,.0f}" if x is not None else f"{'-':>9}" for x in (r["ref_value"], r["loop_value"])]
        log(f"{r['name']:<24}{'/'.join(ym(m / 12) for m in r['best']):>14}{v:>13,.2f}{d:>11,.2f}"
            f"{'-' if tol is None else f'{tol:,.2f}':>7}{'-' if sp is None else f'{sp:,.0f}':>8}  {verdict:<11}"
            f"{costs[0]}{costs[1]}  {res.get('used', '')}")
    log("\ntol: the decided plan's proved tolerance (gap x objective); spread: deepest dip among local-search")
    log("scores (none when every score is proved). (a)/(b) cost: decided objective less the score at that")
    log("search's pick.")
    for r in rows:
        for c in r["close"]:
            log(f"  {r['name']}: best two ages within their tolerances: {c}")
    log(f"\n{'variant':<24}{'federal':>10}{'state':>10}{'payroll':>10}{'Medicare':>10}{'ACA':>10}"
        f"{'SS withheld':>12}{'SS paid':>11}{'bequest':>11}")
    for r in rows:
        m = r["result"].get("metrics")
        if not m:
            continue
        fed = m["federal_income_tax_today"] + m["ltcg_tax_today"] + m["niit_today"]
        log(f"{r['name']:<24}{fed:>10,.0f}{m['state_tax_today']:>10,.0f}{m['payroll_tax_today']:>10,.0f}"
            f"{m['medicare_today']:>10,.0f}{m['aca_today']:>10,.0f}{m['ss_withheld_today']:>12,.0f}"
            f"{m['ss_income_today']:>11,.0f}{m['final_bequest_today']:>11,.0f}")
    _report_netting(rows, log)
    _report_eta(rows, log)
    log("\nScores per sweep (* = not proved: local search's score or an unproved branch-and-bound):")
    for r in rows:
        log(f"  {r['name']}: {len(r['results'])} plans scored")
        for s in r["sweeps"]:
            cells = [(m, v, _cell(r["_cache"], s, m)) for m, v in sorted(s["points"])]
            pts = ", ".join(f"{ym(m / 12)} {v:,.2f}{'' if c.get('proved') else '*'}" for m, v, c in cells)
            log(f"    pass {s['pass']} {s['person']} (held {'/'.join(ym(m / 12) for m in s['held'])}): {pts}")
            for m, dd in s["dips"]:
                kind = "the model's (proved)" if _all_proved(r["_cache"], s, m) else "noise or the model's"
                log(f"      dip at {ym(m / 12)}: {dd:,.2f}, {kind}")


def _report_netting(rows, log):
    """Round-trip netting (#234) in the decided plans, and how many scored plans had any."""
    log("\nRound-trip netting (#234): taxable sale and rebuy the worksheet nets away; today's $.")
    log(f"{'variant':<24}{'netted':>11}{'hidden gain':>13}  years (netted / hidden gain / state tax that year)")
    for r in rows:
        nt = r["result"].get("netting")
        scored = [c for c in r["results"].values() if c.get("netting")]
        some = sum(1 for c in scored if c["netting"]["netted"] > 0.5)
        tail = f"  [{some} of {len(scored)} scored plans net something]"
        if not nt:
            log(f"{r['name']:<24}{'-':>11}{'-':>13}  (not recorded){tail}")
            continue
        yrs = "; ".join(f"{y['year']}: {y['netted']:,.0f} / {y['hidden_gain']:,.0f} / {y['state_tax']:,.0f}"
                        + ("" if y["basis"] else " (no basis)") for y in nt["years"]) or "none"
        log(f"{r['name']:<24}{nt['netted']:>11,.0f}{nt['hidden_gain']:>13,.0f}  {yrs}{tail}")


def _report_eta(rows, log):
    """The deposit-split sensitivity at each variant's decided ages."""
    if not any(r.get("eta") for r in rows):
        return
    log("\nDeposit split (eta) at the decided ages: objective change against the case's split (ages not")
    log("re-decided). resolved = beyond the two plans' proved tolerances.")
    for r in rows:
        if not r.get("eta"):
            continue
        b = r["result"]
        parts = []
        for e, c in r["eta"]:
            if c.get("value") is None or b.get("value") is None:
                parts.append(f"eta {e:.2f}: {c.get('status')}")
                continue
            d = c["value"] - b["value"]
            ta, tb = c.get("tolerance"), b.get("tolerance")
            verdict = ("resolved" if abs(d) > ta + tb else "unresolved") if ta is not None and tb is not None \
                else "not proved"
            nt = (c.get("netting") or {}).get("netted")
            parts.append(f"eta {e:.2f}: {d:+,.2f} ({verdict}{'' if nt is None else f'; netted {nt:,.0f}'})")
        log(f"  {r['name']:<22}{'; '.join(parts)}")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("case")
    ap.add_argument("variants")
    ap.add_argument("--out-dir", default="otherFiles/sweep")
    ap.add_argument("--jobs", type=int, default=1, help="parallel solves in (c) (processes)")
    ap.add_argument("--decider", choices=sorted(DECIDER), default="bb", help="scoring in (c)")
    ap.add_argument("--gap", type=float, default=1e-5, help="relative MIP gap of branch-and-bound in (c)")
    ap.add_argument("--max-time", type=float, default=TIME_LIMIT, help="seconds per MIP solve in (c)")
    ap.add_argument("--passes", type=int, default=2)
    ap.add_argument("--ls-time", type=float, default=7200.0, help="localSearchTime (s); default lifts the cap")
    ap.add_argument("--ls-step-time", type=float, default=1800.0, help="localSearchStepTime (s)")
    ap.add_argument("--no-loop", action="store_true", help="skip stage (b)")
    ap.add_argument("--eta", type=float, nargs="*", default=[0.0, 1.0],
                    help="deposit splits to score at the decided ages (couples; none: --eta with no value)")
    ap.add_argument("--only", nargs="*", help="run only these variants (the first listed is the reference)")
    ap.add_argument("--json", help="write every score to this file")
    ap.add_argument("--report-from", metavar="JSON", help="print the report of a saved --json run, solving nothing")
    ap.add_argument("--objective", default="maxSpending", help="with --report-from: the case's objective")
    args = ap.parse_args()

    def log(msg):
        print(msg, flush=True)

    if args.report_from:
        log(f"work_sweep.py --report-from {args.report_from}, {date.today().isoformat()} (no solves)")
        rows = reload(args.report_from)
        report(rows, args.objective, args.decider, log)
        _summary(rows, args.decider, log)
        return

    with open(args.variants, encoding="utf-8") as f:
        spec = json.load(f)
    variants = build_variants(args.case, spec, args.out_dir)
    if args.only:
        by = {v[0]: v for v in variants}
        variants = [by[n] for n in args.only]
    log(f"work_sweep.py, {date.today().isoformat()}, owlplanner {owl.__version__} from {owl.__file__}")
    log(f"case {args.case}, variants {args.variants}, jobs {args.jobs}, passes {args.passes}, "
        f"decider {DECIDER[args.decider]}"
        + (f" (gap {args.gap:.0e}, relTol {args.gap:.0e}, maxTime {args.max_time:.0f}s per MIP; local search "
           f"for unproved plans)"
           if args.decider == "bb" else "")
        + f", localSearchTime {args.ls_time:.0f}s, localSearchStepTime {args.ls_step_time:.0f}s"
        + f", eta check {', '.join(f'{e:.2f}' for e in args.eta) or 'off'}")
    for name, _, _, sets in variants:
        log(f"  {name}: {'; '.join(sets) or '(no overrides)'}")
    objective = None
    rows = []
    t0 = time.time()
    with ProcessPoolExecutor(max_workers=args.jobs, mp_context=mp.get_context("spawn")) as pool:
        for name, d, dirname, _ in variants:
            if objective is None:
                objective = make_plan(d, dirname)[0].objective
            rows.append(run_variant(name, d, dirname, args, pool, log))
    report(rows, objective, args.decider, log)
    _summary(rows, args.decider, log)
    log(f"Total {time.time() - t0:.0f}s.")
    if args.json:
        with open(args.json, "w", encoding="utf-8") as f:
            json.dump([{k: v for k, v in r.items() if k != "_cache"} for r in rows], f, indent=1, default=str)


def _summary(rows, decider, log):
    def cells(pred):
        return ", ".join(f"{r['name']} {k}" for r in rows for k, c in r["results"].items() if pred(c)) or "none"

    def ls_part(c):
        return c if c.get("used") == "local search" else c.get("fallback", {})

    log("")
    if decider == "bb":
        log(f"Scored plans not proved by branch-and-bound: {cells(lambda c: not c.get('proved'))}")
        log(f"  of which a MIP stopped on time: {cells(lambda c: (c.get('bb') or c).get('time_limit_runs'))}")
        log(f"  of which a MIP stopped at the node limit: {cells(lambda c: (c.get('bb') or c).get('node_limit_runs'))}")
        log(f"  of which local search's score was kept: {cells(lambda c: c.get('used') == 'local search')}")
    log(f"Local-search scores where its default caps (300 s / 60 s) would have bound: "
        f"{cells(lambda c: ls_part(c).get('default_caps_bind'))}")
    log(f"Local-search scores that spent the raised time budget: {cells(lambda c: ls_part(c).get('budget_spent'))}")
    log(f"Scored plans not solved: {cells(lambda c: c['status'] != 'solved')}")
    if decider == "bb":
        gaps = [c["gap"] for r in rows for c in r["results"].values() if c.get("proved")]
        if gaps:
            log(f"Reported gaps of proved plans: {min(gaps):.1e} to {max(gaps):.1e}")


if __name__ == "__main__":
    main()
