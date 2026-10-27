"""Generate an empty HFP workbook for the Phase 0 baseline case.

Run from the repository root:

    uv run python fork-notes/phase0/gen_hfp_us.py                          # otherFiles/HFP_us.xlsx
    uv run python fork-notes/phase0/gen_hfp_us.py otherFiles/HFP_us_2027.xlsx

It refuses to overwrite an existing workbook, so a filled-in one is never lost.

Columns per person sheet (exact headers, order may vary; only 'year' is required):
  year, anticipated wages, months worked, other inc, net inv,
  taxable ctrb, 401k ctrb, employer ctrb, Roth 401k ctrb, IRA ctrb, Roth IRA ctrb, HSA ctrb,
  Roth conv, Roth conv fixed, QCD, big-ticket items
(upstream's order, hfp_io._timeHorizonItems)

'anticipated wages' are the wages in AGI: pay minus pre-tax 401k, deductible IRA and payroll
HSA contributions (upstream #221). An after-tax contribution paid from pay (Roth 401k, Roth IRA,
taxable) stays in the wages and is paid by a negative 'big-ticket items' entry of the same amount.
'401k ctrb' is the employee's own pre-tax deferral; an employer match or nonelective contribution
goes in 'employer ctrb' (same account, not earnings: no payroll tax, no earnings test; upstream
#229/#230). A 401k deferral is earnings even in a year with no wages; 'IRA ctrb' counts as
earnings only in a year with wages or a 401k deferral (otherwise it is taken as a spousal IRA).
'months worked' (0-12) is read by the SSA earnings test: enter it in the year a spouse stops
working (the grace year); 0 or blank means 12 in a year with wages.
'big-ticket items' (negative = outflow) and 'other inc' (negative = an above-the-line deduction,
e.g. an HSA contribution without wages, upstream #225) are the only signed columns.

The Budget sheet starts with three lines (the template's spending_profile is "budget"):
'core' (discretionary, from the first year to the plan end, survivor blank = the case's 60%):
TODO amount, today's spending other than the lines below. Under maxSpending its amount only sets
the shape (the solver scales it, so the result is the same for any positive amount); under
maxBequest, or with netSpending, it is the spending: enter the real figure.
'gold over benchmark' (medical, essential): the cheapest gold plan's premium over the silver
benchmark for the couple, 2027 NY Rating Region 4 approved rates (Fidelis gold 25303NY0030001
2,438.10/mo vs Anthem silver 41046NY0010004 2,107.36/mo, x 12 = 3,969; fork-notes/aca-ny/README.md,
"Gold plans"). Owl prices the marketplace at the benchmark, so this is the extra for buying gold.
2027-2028: both on the marketplace (aca_start_year 2027); survivor 50, one adult pays half the
couple's rate. 2029: the template's dates (A born in January, B in December 1964) leave B alone on
the marketplace for about 11 months, so half the couple's extra x 11/12 = 1,819; survivor 100.
Redo both lines when the real dates and aca_start_year are in. Amounts in 2027 dollars entered
as today's (2026): about one year of inflation high.
"""

import sys
from pathlib import Path

import openpyxl

NAMES = ("SpouseA", "SpouseB")  # must match basic_info.names in the case file
YEARS = list(range(2026, 2057))  # through the year both reach 92 (born 1964)
COLUMNS = [
    "year", "anticipated wages", "months worked", "other inc", "net inv",
    "taxable ctrb", "401k ctrb", "employer ctrb", "Roth 401k ctrb", "IRA ctrb", "Roth IRA ctrb", "HSA ctrb",
    "Roth conv", "Roth conv fixed", "QCD", "big-ticket items",
]
# Household sheets: headers, and BUDGET_LINES in Budget (more rows per scenario; phase0-scenarios.md, section 5).
HOUSE_SHEETS = {
    "Debts": ["active", "name", "type", "year", "term", "amount", "rate", "property"],
    "Fixed Assets": ["active", "name", "type", "year", "basis", "value", "rate", "yod", "commission"],
    "Budget": ["active", "name", "type", "year", "end", "amount", "rate", "survivor", "essential"],
}
# Lines written into the Budget sheet (see the docstring).
BUDGET_LINES = [
    [True, "core", "core", YEARS[0], 0, 100000, 0, None, False],  # TODO amount (see the docstring)
    [True, "gold over benchmark", "medical", 2027, 2028, 3969, 0, 50, True],
    [True, "gold over benchmark, B alone", "medical", 2029, 2029, 1819, 0, 100, True],
]

out = Path(sys.argv[1] if len(sys.argv) > 1 else "otherFiles/HFP_us.xlsx")
if out.exists():
    sys.exit(f"{out} exists; not overwriting it. Pass another path or delete it first.")
out.parent.mkdir(parents=True, exist_ok=True)

wb = openpyxl.Workbook()
wb.remove(wb.active)
for name in NAMES:
    ws = wb.create_sheet(name)
    ws.append(COLUMNS)
    for year in YEARS:
        ws.append([year] + [0] * (len(COLUMNS) - 1))
for sheet, columns in HOUSE_SHEETS.items():
    ws = wb.create_sheet(sheet)
    ws.append(columns)
    if sheet == "Budget":
        for line in BUDGET_LINES:
            ws.append(line)
wb.save(out)

print(f"Wrote {out}: sheets {wb.sheetnames}, years {YEARS[0]}-{YEARS[-1]} ({len(YEARS)} rows)")
print("Fill in: anticipated wages (as in AGI, upstream #221) and months worked (the year work stops) for the working years;")
print("         big-ticket items, negative, for long-term-care and other one-off outflows.")
print("         Budget sheet (spending_profile = \"budget\"): set 'core' (TODO amount), keep or drop the gold lines; add")
print("         rent / property tax / insurance /")
print("         maintenance / car / travel / care lines (year, end, amount in today's $, rate, survivor %).")
print("         Buying: Fixed Assets (residence) and Debts (mortgage; property = the residence's name).")
