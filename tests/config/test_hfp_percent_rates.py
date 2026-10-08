"""Rates in the Debts and Fixed Assets sheets: a percent-formatted cell is a fraction, any other is percent."""

import io

import numpy as np
import openpyxl
import pandas as pd
import pytest

import owlplanner as owl
from owlplanner.hfp_io import conditionDebtsAndFixedAssetsDF as cond


def _plan():
    p = owl.Plan(["Joe", "Jane"], ["1964-03-15", "1964-09-15"], [89, 92], "rates", verbose=False)
    y = int(p.year_n[0])
    p.houseLists["Fixed Assets"] = cond(pd.DataFrame([dict(active=True, name="home", type="residence", year=y,
                                                           basis=500000.0, value=800000.0, rate=0.5, yod=0,
                                                           commission=6.0)]), "Fixed Assets")
    p.houseLists["Debts"] = cond(pd.DataFrame([dict(active=True, name="mtg", type="mortgage", year=y, term=30,
                                                    amount=400000.0, rate=0.9)]), "Debts")
    return p


def _read(src):
    q = owl.Plan(["Joe", "Jane"], ["1964-03-15", "1964-09-15"], [89, 92], "rates", verbose=False)
    q.readHFP(src)
    return q


def test_rates_below_one_percent_survive_a_round_trip(tmp_path):
    """0.5% real growth and a 0.9% loan were read back as 50% and 90%."""
    fname = str(tmp_path / "HFP_rates.xlsx")
    _plan().saveHFP(fname, overwrite=True)
    q = _read(fname)
    assert q.houseLists["Fixed Assets"]["rate"].iloc[0] == pytest.approx(0.5)
    assert q.houseLists["Debts"]["rate"].iloc[0] == pytest.approx(0.9)


def _percent_format(fname, sheet, column, value):
    wb = openpyxl.load_workbook(fname)
    ws = wb[sheet]
    k = [c.value for c in ws[1]].index(column)
    cell = ws.cell(row=2, column=k + 1)
    cell.value = value
    cell.number_format = "0.00%"
    wb.save(fname)


def test_percent_formatted_cells_are_fractions(tmp_path):
    fname = str(tmp_path / "HFP_pct.xlsx")
    _plan().saveHFP(fname, overwrite=True)
    _percent_format(fname, "Debts", "rate", 0.045)  # shown as 4.50%
    _percent_format(fname, "Fixed Assets", "commission", 0.06)  # shown as 6.00%
    q = _read(fname)
    assert q.houseLists["Debts"]["rate"].iloc[0] == pytest.approx(4.5)
    assert q.houseLists["Fixed Assets"]["commission"].iloc[0] == pytest.approx(6.0)
    assert q.houseLists["Fixed Assets"]["rate"].iloc[0] == pytest.approx(0.5)  # not formatted: as typed


def test_uploaded_stream_is_read_and_rewound(tmp_path):
    """The UI hands over an in-memory upload."""
    fname = str(tmp_path / "HFP_stream.xlsx")
    _plan().saveHFP(fname, overwrite=True)
    _percent_format(fname, "Debts", "rate", 0.045)
    with open(fname, "rb") as f:
        buf = io.BytesIO(f.read())
    q = _read(buf)
    assert q.houseLists["Debts"]["rate"].iloc[0] == pytest.approx(4.5)
    assert q.houseLists["Fixed Assets"]["rate"].iloc[0] == pytest.approx(0.5)
    assert buf.tell() == 0


def test_tables_given_as_dataframes_are_read_as_typed():
    from owlplanner import hfp_io

    log = type("Log", (), {"vprint": lambda self, *a, **k: None, "print": lambda self, *a, **k: None})()
    p = _plan()
    house = hfp_io._conditionHouseTables({"Debts": p.houseLists["Debts"]}, log)
    assert house["Debts"]["rate"].iloc[0] == pytest.approx(0.9)
    assert np.isclose(house["Debts"]["amount"].iloc[0], 400000.0)
