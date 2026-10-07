"""Tests for the optional `payoff` column of the Debts sheet: the balance is paid in that year."""

import numpy as np
import pandas as pd
import pytest

import owlplanner as owl
from owlplanner import debts
from owlplanner.hfp_io import conditionDebtsAndFixedAssetsDF

THISYEAR = owl.Plan(["A"], ["1960-01-01"], [80], "t", verbose=False).year_n[0]


def _loan(payoff=None, year=2020, term=30, amount=400_000.0, rate=6.0):
    row = {"active": True, "name": "mtg", "type": "mortgage", "year": year, "term": term,
           "amount": amount, "rate": rate}
    if payoff is not None:
        row["payoff"] = payoff
    return pd.DataFrame([row])


class TestPayoffSchedule:
    def test_no_column_is_unchanged(self):
        """A table without the column (an older workbook) runs to term."""
        df = _loan()
        pay = debts.get_debt_payments_array(df, 40, 2020)
        assert np.count_nonzero(pay) == 30
        assert np.allclose(pay[:30], debts.calculate_annual_payment(400_000.0, 6.0, 30))

    def test_zero_is_no_payoff(self):
        assert np.allclose(debts.get_debt_payments_array(_loan(0), 40, 2020),
                           debts.get_debt_payments_array(_loan(), 40, 2020))

    def test_payoff_pays_the_balance_and_stops(self):
        df = _loan(2030)
        pay = debts.get_debt_payments_array(df, 40, 2020)
        annual = debts.calculate_annual_payment(400_000.0, 6.0, 30)
        balance = debts.calculate_remaining_balance(400_000.0, 6.0, 30, 10)
        assert np.allclose(pay[:10], annual)
        assert pay[10] == pytest.approx(balance)
        assert np.all(pay[11:] == 0)

    def test_payoff_matches_the_balance_at_the_start_of_that_year(self):
        df = _loan(2030)
        bal = debts.get_debt_balances_array(df, 40, 2020)
        pay = debts.get_debt_payments_array(df, 40, 2020)
        assert bal[10] == pytest.approx(pay[10])  # owed at the start of 2030, paid in 2030
        assert np.all(bal[11:] == 0)
        assert debts.get_debt_balances_for_year(df, 2029) == pytest.approx(bal[10])
        assert debts.get_debt_balances_for_year(df, 2030) == 0.0

    def test_payoff_before_plan_end_leaves_no_balance(self):
        assert debts.get_remaining_debt_balance(_loan(2030), 15, 2020) == 0.0
        # A payoff after the plan end leaves the same balance as no payoff.
        assert debts.get_remaining_debt_balance(_loan(2040), 15, 2020) == pytest.approx(
            debts.get_remaining_debt_balance(_loan(), 15, 2020))

    def test_payoff_at_or_after_term_changes_nothing(self):
        for payoff in (2050, 2060):
            assert np.allclose(debts.get_debt_payments_array(_loan(payoff), 40, 2020),
                               debts.get_debt_payments_array(_loan(), 40, 2020))

    def test_payoff_in_the_start_year_pays_the_principal(self):
        pay = debts.get_debt_payments_array(_loan(2020), 5, 2020)
        assert pay[0] == pytest.approx(400_000.0) and np.all(pay[1:] == 0)

    def test_payments_for_year(self):
        df = _loan(2030)
        assert debts.get_debt_payments_for_year(df, 2030) == pytest.approx(
            debts.calculate_remaining_balance(400_000.0, 6.0, 30, 10))
        assert debts.get_debt_payments_for_year(df, 2031) == 0.0


def test_condition_fills_a_missing_payoff_with_zero():
    df = conditionDebtsAndFixedAssetsDF(_loan(), "Debts")
    assert "payoff" in df.columns and df["payoff"].iloc[0] == 0


def _plan(payoff):
    p = owl.Plan(["Joe", "Jane"], ["1964-03-15", "1965-09-15"], [89, 92], "payoff", verbose=False)
    p.setSpendingProfile("flat")
    p.setAccountBalances(taxable=[600, 600], taxDeferred=[0, 0], taxFree=[0, 0])
    p.setAllocationRatios("individual", generic=np.array([[[60, 40, 0, 0], [60, 40, 0, 0]]] * 2))
    p.setRates("conservative")
    p.setSocialSecurity([2800, 2200], [70, 70])
    p.houseLists["Debts"] = conditionDebtsAndFixedAssetsDF(
        _loan(payoff, year=THISYEAR, term=30, amount=500_000.0, rate=6.5), "Debts")
    return p


def test_plan_pays_off_and_reports_no_debt_after():
    p = _plan(THISYEAR + 8)
    p.solve("maxSpending", options={"bequest": 0, "withMedicare": "None", "withSSTaxability": 0.85})
    assert p.caseStatus == "solved"
    assert p.debt_payments_n[8] == pytest.approx(debts.calculate_remaining_balance(500_000.0, 6.5, 30, 8))
    assert np.all(p.debt_payments_n[9:] == 0)
    assert p.remaining_debt_balance == 0.0
    assert np.all(p.fixed_assets_debt_balances_remaining_n[9:] == 0)


def test_hfp_round_trip_keeps_payoff(tmp_path):
    p1 = _plan(THISYEAR + 8)
    fname = str(tmp_path / "HFP_payoff.xlsx")
    p1.saveHFP(fname, overwrite=True)
    p2 = _plan(None)
    p2.readHFP(fname)
    assert int(p2.houseLists["Debts"]["payoff"].iloc[0]) == THISYEAR + 8


def test_hfp_without_payoff_column_loads(tmp_path):
    """A workbook written before the column existed loads, with no payoff."""
    p1 = _plan(None)
    fname = str(tmp_path / "HFP_old.xlsx")
    p1.saveHFP(fname, overwrite=True)
    from openpyxl import load_workbook

    wb = load_workbook(fname)
    ws = wb["Debts"]
    header = [c.value for c in ws[1]]
    ws.delete_cols(header.index("payoff") + 1)
    wb.save(fname)
    p2 = _plan(None)
    p2.readHFP(fname)
    assert int(p2.houseLists["Debts"]["payoff"].iloc[0]) == 0
    assert float(p2.houseLists["Debts"]["amount"].iloc[0]) == pytest.approx(500_000.0)
